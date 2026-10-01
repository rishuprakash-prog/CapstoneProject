"""Classroom Transaction (CRT) report: weekly school-health sessions by theme.

The CRT export has one row per *activity*. A session (one class-section, one date, one
theme/module) is repeated once for every activity done in it, each repeat carrying the same
boys/girls counts. Summing raw rows would double-count participants, so everything here is
computed on de-duplicated sessions.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass

import pandas as pd

# role -> accepted normalised column names (first match wins)
CRT_COLUMNS = {
    "district": ["district"],
    "block": ["block"],
    "school": ["udise", "udise_code", "school_id"],
    "school_name": ["schoolname", "school_name"],
    "class": ["class"],
    "section": ["section"],
    "boys": ["boysp", "boys", "boys_present"],
    "girls": ["girlsp", "girls", "girls_present"],
    "other": ["otherp", "other", "others"],
    "fy": ["finyear", "fin_year", "financial_year", "fy"],
    "week": ["weak", "week", "week_no"],
    "date": ["sessiondate", "session_date"],
    "theme": ["module", "theme"],
    "activity": ["activity"],
}
REQUIRED = ["district", "school", "boys", "girls", "theme"]
SESSION_KEY = ["school", "class", "section", "date", "theme"]
OUTLIER_LIMIT = 200  # participants in one class-section session above this are implausible


def _norm(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(name).strip().lower()).strip("_")


def map_columns(df: pd.DataFrame) -> dict[str, str]:
    normed = {_norm(c): c for c in df.columns}
    found = {}
    for role, names in CRT_COLUMNS.items():
        for n in names:
            if n in normed:
                found[role] = normed[n]
                break
    return found


def is_crt(df: pd.DataFrame) -> bool:
    return all(r in map_columns(df) for r in REQUIRED)


@dataclass
class CRTData:
    sessions: pd.DataFrame  # one row per session, standard column names
    raw_rows: int
    duplicate_rows: int

    @property
    def fin_years(self) -> list[str]:
        return sorted(self.sessions["fy"].dropna().unique().tolist())


def prepare_sessions(df: pd.DataFrame) -> CRTData:
    cols = map_columns(df)
    std = pd.DataFrame({role: df[col] for role, col in cols.items()})
    for role in ("class", "section", "date", "fy", "week", "block", "school_name", "activity"):
        if role not in std:
            std[role] = "" if role != "week" else 0
    for role in ("boys", "girls", "other"):
        std[role] = pd.to_numeric(std.get(role, 0), errors="coerce").fillna(0)
    std["district"] = std["district"].astype(str).str.strip().str.title()
    std["block"] = std["block"].astype(str).str.strip().str.title()
    std["theme"] = std["theme"].astype(str).str.replace(r"\s+", " ", regex=True).str.strip()
    std["fy"] = std["fy"].astype(str).str.strip()
    std["week"] = pd.to_numeric(std["week"], errors="coerce").fillna(0).astype(int)

    dupes = int(df.duplicated().sum())
    std = std[~df.duplicated()]
    sessions = (std.groupby(SESSION_KEY, dropna=False)
                .agg(district=("district", "first"), block=("block", "first"),
                     school_name=("school_name", "first"), fy=("fy", "first"), week=("week", "first"),
                     boys=("boys", "max"), girls=("girls", "max"), other=("other", "max"),
                     activities=("activity", "nunique"))
                .reset_index())
    sessions["total"] = sessions["boys"] + sessions["girls"] + sessions["other"]
    return CRTData(sessions, len(df), dupes)


def filter_sessions(data: CRTData, fys: list[str] | None = None, districts: list[str] | None = None,
                    classes: list[str] | None = None, drop_outliers: bool = False) -> pd.DataFrame:
    s = data.sessions
    if fys:
        s = s[s["fy"].isin(fys)]
    if districts:
        s = s[s["district"].isin(districts)]
    if classes:
        s = s[s["class"].isin(classes)]
    if drop_outliers:
        s = s[(s["total"] > 0) & (s["total"] <= OUTLIER_LIMIT)]
    return s


# ---------------------------------------------------------------- report tables

def _gender_cols(t: pd.DataFrame) -> pd.DataFrame:
    t["Total Participants"] = t["Boys"] + t["Girls"]
    t["Girls %"] = (100 * t["Girls"] / t["Total Participants"].where(t["Total Participants"] > 0)).round(1)
    return t


def fy_summary(s: pd.DataFrame) -> pd.DataFrame:
    t = s.groupby("fy").agg(**{
        "Districts": ("district", "nunique"), "Blocks": ("block", "nunique"),
        "Schools Reported": ("school", "nunique"), "Sessions": ("school", "size"),
        "Weeks Reported": ("week", "nunique"), "Boys": ("boys", "sum"), "Girls": ("girls", "sum"),
    })
    t.index.name = "Financial Year"
    return _gender_cols(t)


def district_report(s: pd.DataFrame) -> pd.DataFrame:
    """Schools reported and boys/girls participation per district, with a state total row."""
    def summarise(g: pd.DataFrame) -> dict:
        return {
            "Blocks": g["block"].nunique(), "Schools Reported": g["school"].nunique(),
            "Sessions": len(g), "Boys": g["boys"].sum(), "Girls": g["girls"].sum(),
            "Avg Boys / Session": round(g["boys"].mean(), 1), "Avg Girls / Session": round(g["girls"].mean(), 1),
            "Avg Sessions / School": round(len(g) / max(g["school"].nunique(), 1), 1),
        }
    t = pd.DataFrame({d: summarise(g) for d, g in s.groupby("district")}).T
    t = t.sort_values("Schools Reported", ascending=False)
    if len(s):
        t.loc["STATE TOTAL"] = summarise(s)
    t.index.name = "District"
    t = _gender_cols(t)
    ints = ["Blocks", "Schools Reported", "Sessions", "Boys", "Girls", "Total Participants"]
    t[ints] = t[ints].astype(int)
    return t


def theme_report(s: pd.DataFrame) -> pd.DataFrame:
    t = s.groupby("theme").agg(**{
        "Schools": ("school", "nunique"), "Sessions": ("school", "size"),
        "Boys": ("boys", "sum"), "Girls": ("girls", "sum"),
        "Avg Boys / Session": ("boys", "mean"), "Avg Girls / Session": ("girls", "mean"),
    }).round(1).sort_values("Sessions", ascending=False)
    t.index.name = "Theme (Module)"
    return _gender_cols(t)


def weekly_report(s: pd.DataFrame) -> pd.DataFrame:
    t = s.groupby(["fy", "week"]).agg(**{
        "Week Start": ("date", "min"), "Schools": ("school", "nunique"), "Sessions": ("school", "size"),
        "Boys": ("boys", "sum"), "Girls": ("girls", "sum"),
    }).reset_index().rename(columns={"fy": "Financial Year", "week": "Week"})
    t["Week Start"] = pd.to_datetime(t["Week Start"], errors="coerce").dt.strftime("%d-%b-%Y")
    return _gender_cols(t)


def weekly_theme(s: pd.DataFrame) -> pd.DataFrame:
    """Long table: week x theme x gender participation (for the weekly theme chart)."""
    t = s.groupby(["week", "theme"])[["boys", "girls"]].sum().reset_index()
    return t.rename(columns={"week": "Week", "theme": "Theme", "boys": "Boys", "girls": "Girls"})


def district_theme_matrix(s: pd.DataFrame, value: str = "sessions") -> pd.DataFrame:
    if value == "sessions":
        m = s.pivot_table(index="district", columns="theme", values="school", aggfunc="size", fill_value=0)
    else:
        s = s.assign(participants=s["boys"] + s["girls"])
        m = s.pivot_table(index="district", columns="theme", values="participants", aggfunc="sum", fill_value=0)
    m.index.name, m.columns.name = "District", None
    return m.loc[m.sum(axis=1).sort_values(ascending=False).index]


def data_quality(data: CRTData, s: pd.DataFrame) -> pd.DataFrame:
    rows = {
        "Raw rows in file (one per activity)": data.raw_rows,
        "Exact duplicate rows removed": data.duplicate_rows,
        "Unique sessions after de-duplication": len(data.sessions),
        "Sessions in current filter": len(s),
        "Sessions with 0 participants": int((s["total"] == 0).sum()),
        f"Sessions with > {OUTLIER_LIMIT} participants (likely entry error)": int((s["total"] > OUTLIER_LIMIT).sum()),
    }
    return pd.DataFrame({"Check": rows.keys(), "Value": rows.values()})


# ---------------------------------------------------------------- AI context & export

def build_context(data: CRTData, s: pd.DataFrame, filters: str) -> str:
    parts = [
        "## Dataset: Classroom Transaction (weekly school health & wellness sessions) report",
        "Unit of analysis = session (one class-section, one date, one theme). Boys/Girls = participants "
        "present in that session. 'Avg / Session' = mean participants per session. Participation totals "
        "count attendances, not unique students.",
        f"Filters applied: {filters}",
        "## Financial year summary", fy_summary(s).to_csv(),
        "## District report", district_report(s).to_csv(),
        "## Theme-wise participation", theme_report(s).to_csv(),
        "## Weekly participation", weekly_report(s).to_csv(index=False),
        "## District x theme (sessions)", district_theme_matrix(s).to_csv(),
        "## Data quality", data_quality(data, s).to_csv(index=False),
    ]
    return "\n".join(parts)


def excel_report(sheets: dict[str, pd.DataFrame], title: str) -> bytes:
    """Multi-sheet Excel report with native Excel charts next to each table."""
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name[:31], startrow=2)
        wb = xw.book
        head_fill = PatternFill("solid", fgColor="1F4E79")
        for name, df in sheets.items():
            ws = wb[name[:31]]
            ws["A1"] = f"{title} — {name}"
            ws["A1"].font = Font(bold=True, size=13, color="1F4E79")
            for cell in ws[3]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = head_fill
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            for col in ws.columns:
                width = max(len(str(c.value or "")) for c in col[2:60])
                ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 45)
            ws.freeze_panes = "B4"

        def add_bar(sheet, cols, chart_title, anchor, n_rows, horizontal=False, stacked=False):
            ws = wb[sheet]
            df = sheets[sheet]
            ch = BarChart()
            ch.type = "bar" if horizontal else "col"
            if stacked:
                ch.grouping, ch.overlap = "stacked", 100
            ch.title, ch.height, ch.width = chart_title, 9 + 0.25 * n_rows, 20
            for c in cols:
                idx = list(df.columns).index(c) + 2
                ch.add_data(Reference(ws, min_col=idx, min_row=3, max_row=3 + n_rows), titles_from_data=True)
            ch.set_categories(Reference(ws, min_col=1, min_row=4, max_row=3 + n_rows))
            ws.add_chart(ch, anchor)

        if "District Report" in sheets:
            n = len(sheets["District Report"]) - 1  # exclude state total row
            col = chr(ord("A") + len(sheets["District Report"].columns) + 2)
            add_bar("District Report", ["Schools Reported"], "Schools reported per district", f"{col}3", n, True)
            add_bar("District Report", ["Avg Boys / Session", "Avg Girls / Session"],
                    "Average boys & girls per session", f"{col}{30 + n // 2}", n, True)
        if "Theme Report" in sheets:
            n = len(sheets["Theme Report"])
            col = chr(ord("A") + len(sheets["Theme Report"].columns) + 2)
            add_bar("Theme Report", ["Boys", "Girls"], "Theme-wise boys & girls participation", f"{col}3", n, True, True)
        if "Weekly Report" in sheets:
            ws, df = wb["Weekly Report"], sheets["Weekly Report"]
            n = len(df)
            ch = LineChart()
            ch.title, ch.height, ch.width = "Weekly participation", 9, 22
            for c in ["Boys", "Girls"]:
                idx = list(df.columns).index(c) + 2
                ch.add_data(Reference(ws, min_col=idx, min_row=3, max_row=3 + n), titles_from_data=True)
            wk = list(df.columns).index("Week") + 2
            ch.set_categories(Reference(ws, min_col=wk, min_row=4, max_row=3 + n))
            ws.add_chart(ch, f"{chr(ord('A') + len(df.columns) + 2)}3")
    return buf.getvalue()
